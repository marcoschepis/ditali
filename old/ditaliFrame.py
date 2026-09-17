import openpyxl
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Patch
import os

# --- CONFIGURAZIONE PERCORSI ---
EXCEL_FILE = "ditali.xlsx"
FRAME_DRIVE = "F:/PHOTO"

def hex_from_cell(cell):
    """Estrae il colore HEX di riempimento dalla cella Excel (se presente)."""
    if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
        rgb = str(cell.fill.start_color.rgb)
        if len(rgb) == 8:    # Formato ARGB -> rimuove canale Alfa
            return "#" + rgb[2:].upper()
        elif len(rgb) == 6:  # Formato RGB
            return "#" + rgb.upper()
    return None

def genera_statistiche_globali(wb, sheet_names, Mappa_legenda_globale, totale_globale):
    """Genera una scheda grafica con le statistiche dettagliate per ogni categoria."""
    conteggio_categorie = {color: 0 for color in Mappa_legenda_globale}
    ditali_senza_categoria = 0

    # Conteggio ditali per ogni categoria/colore
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                val = cell.value
                hex_col = hex_from_cell(cell)

                is_num = False
                num_val = 0
                if isinstance(val, (int, float)):
                    num_val = val
                    is_num = True
                elif isinstance(val, str) and val.strip().replace('.', '', 1).replace(',', '', 1).isdigit():
                    num_val = float(val)
                    is_num = True

                if is_num:
                    if hex_col in conteggio_categorie:
                        conteggio_categorie[hex_col] += int(num_val)
                    else:
                        ditali_senza_categoria += int(num_val)

    if ditali_senza_categoria > 0:
        conteggio_categorie['#334155'] = ditali_senza_categoria
        Mappa_legenda_globale['#334155'] = 'Senza Categoria'

    # Creazione Dashboard Statistiche
    fig, ax = plt.subplots(figsize=(10, 7), dpi=100)
    fig.patch.set_facecolor('#0f172a')
    ax.set_facecolor('#0f172a')

    # Intestazione
    fig.suptitle(f" STATISTICHE CATEGORIE  |  TOTALE GLOBALE: {int(totale_globale)} DITALI ", 
                 color='#38bdf8', fontsize=14, fontweight='bold', y=0.95,
                 bbox=dict(boxstyle='round,pad=0.5', facecolor='#1e293b', edgecolor='#0284c7', linewidth=1.5))

    labels = [Mappa_legenda_globale[color] for color in conteggio_categorie]
    values = [conteggio_categorie[color] for color in conteggio_categorie]
    colors = [color for color in conteggio_categorie]

    # Grafico a Barre Orizzontali
    y_pos = range(len(labels))
    bars = ax.barh(y_pos, values, color=colors, edgecolor='#64748b', height=0.55)

    ax.set_yticks(y_pos)
    ax.set_yticklabels(labels, color='#f8fafc', fontsize=11, fontweight='bold')
    ax.invert_yaxis()  # Dall'alto verso il basso
    ax.tick_params(axis='x', colors='#94a3b8')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#475569')
    ax.spines['bottom'].set_color('#475569')

    # Aggiunge il conteggio numerico e la percentuale sopra/accanto a ogni barra
    for bar in bars:
        width = bar.get_width()
        percentuale = (width / totale_globale * 100) if totale_globale > 0 else 0
        ax.text(width + (max(values)*0.02 if values else 1), bar.get_y() + bar.get_height()/2, 
                f"{int(width)} ditali ({percentuale:.1f}%)", 
                va='center', color='#38bdf8', fontsize=10, fontweight='bold')

    plt.tight_layout(rect=[0.05, 0.05, 0.95, 0.90])

    # Salvataggio foto statistiche
    local_stat_img = "display_Statistiche.jpg"
    plt.savefig(local_stat_img, facecolor=fig.get_facecolor(), edgecolor='none', bbox_inches='tight')
    print(f"✓ Immagine statistiche salvata: {os.path.abspath(local_stat_img)}")

    frame_stat_img = os.path.join(FRAME_DRIVE, "ditali_Statistiche.jpg")
    try:
        plt.savefig(frame_stat_img, facecolor=fig.get_facecolor(), edgecolor='none', bbox_inches='tight')
        print(f"✓ Copia cornice aggiornata: {frame_stat_img}")
    except Exception:
        pass

    plt.close()

def genera_mappe_stanze_singole():
    if not os.path.exists(EXCEL_FILE):
        print(f"Errore: Il file {EXCEL_FILE} non esiste nella cartella corrente!")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    sheet_names = wb.sheetnames

    # 1. Calcolo del TOTALE GLOBALE
    totale_globale = 0
    mappa_legenda_globale = {}

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        
        # Mappa i colori presi dalla riga 1
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=c)
            val = cell.value
            hex_col = hex_from_cell(cell)
            if hex_col and val is not None and str(val).strip():
                mappa_legenda_globale[hex_col] = str(val).strip()

        # Somma i numeri dalla riga 2 in poi
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(row=r, column=c).value
                if isinstance(val, (int, float)):
                    totale_globale += val
                elif isinstance(val, str) and val.strip().replace('.', '', 1).replace(',', '', 1).isdigit():
                    totale_globale += float(val)

    # 2. Genera l'immagine delle Statistiche generali
    genera_statistiche_globali(wb, sheet_names, mappa_legenda_globale, totale_globale)

    # 3. Genera l'immagine per ciascuna stanza
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        max_row = ws.max_row
        max_col = ws.max_column
        merged_ranges = ws.merged_cells.ranges

        totale_stanza = 0
        colori_legenda = {}

        for c in range(1, max_col + 1):
            cell = ws.cell(row=1, column=c)
            val = cell.value
            hex_col = hex_from_cell(cell)
            if hex_col and val is not None and str(val).strip():
                colori_legenda[hex_col] = str(val).strip()

        for r in range(2, max_row + 1):
            for c in range(1, max_col + 1):
                val = ws.cell(row=r, column=c).value
                if isinstance(val, (int, float)):
                    totale_stanza += val
                elif isinstance(val, str) and val.strip().replace('.', '', 1).replace(',', '', 1).isdigit():
                    totale_stanza += float(val)

        fig, ax = plt.subplots(figsize=(10, 7), dpi=100)
        fig.patch.set_facecolor('#0f172a')
        ax.set_facecolor('#0f172a')
        ax.axis('off')

        titolo_completo = f" STANZA: {sheet_name.upper()}  |  TOTALE STANZA: {int(totale_stanza)}  |  TOTALE GLOBALE: {int(totale_globale)} "
        fig.suptitle(titolo_completo, color='#38bdf8', fontsize=13, fontweight='bold', y=0.95,
                     bbox=dict(boxstyle='round,pad=0.5', facecolor='#1e293b', edgecolor='#0284c7', linewidth=1.5))

        num_righe_visibili = max(1, max_row - 1)
        bbox_x0, bbox_y0, bbox_w, bbox_h = 0.05, 0.12, 0.90, 0.76
        
        frame_border = FancyBboxPatch((bbox_x0 - 0.01, bbox_y0 - 0.01), bbox_w + 0.02, bbox_h + 0.02,
                                     boxstyle="round,pad=0.01,rounding_size=0.02",
                                     facecolor='#1e293b', edgecolor='#0284c7', linewidth=2.0,
                                     transform=ax.transAxes)
        ax.add_patch(frame_border)

        col_width = bbox_w / max_col
        row_height = bbox_h / num_righe_visibili
        skip_cells = set()

        def draw_styled_cell(rx, ry, rw, rh, val, cell_color=None):
            is_numeric = isinstance(val, (int, float)) or (isinstance(val, str) and val.strip().replace('.', '', 1).replace(',', '', 1).isdigit())

            if is_numeric:
                bg_color = cell_color if cell_color else '#334155'
                padding = 0.004
                cell_box = FancyBboxPatch((rx + padding, ry + padding), rw - 2*padding, rh - 2*padding,
                                          boxstyle="round,pad=0.004,rounding_size=0.012",
                                          facecolor=bg_color, edgecolor='#64748b', linewidth=1.2,
                                          transform=ax.transAxes)
                ax.add_patch(cell_box)

                text_val = str(int(float(val))) if isinstance(val, (int, float)) or (isinstance(val, str) and val.replace('.', '', 1).isdigit()) else str(val)
                ax.text(rx + rw/2, ry + rh/2, text_val, color='#f8fafc', fontsize=13, fontweight='bold',
                        ha='center', va='center', transform=ax.transAxes)
            else:
                text_val = str(val).strip()
                if text_val:
                    ax.text(rx + rw/2, ry + rh/2, text_val, color='#38bdf8', fontsize=11, fontweight='bold',
                            ha='center', va='center', transform=ax.transAxes)

        for m_range in merged_ranges:
            min_col, min_row, max_col_m, max_row_m = m_range.min_col, m_range.min_row, m_range.max_col, m_range.max_row
            if min_row < 2:
                continue

            top_cell = ws.cell(row=min_row, column=min_col)
            val = top_cell.value
            cell_color = hex_from_cell(top_cell)

            for r in range(min_row, max_row_m + 1):
                for c in range(min_col, max_col_m + 1):
                    skip_cells.add((r, c))

            if val is None or str(val).strip() == "":
                continue

            rx = bbox_x0 + (min_col - 1) * col_width
            ry = bbox_y0 + (max_row - max_row_m) * row_height
            rw = (max_col_m - min_col + 1) * col_width
            rh = (max_row_m - min_row + 1) * row_height

            draw_styled_cell(rx, ry, rw, rh, val, cell_color)

        for r in range(2, max_row + 1):
            for c in range(1, max_col + 1):
                if (r, c) in skip_cells:
                    continue
                cell = ws.cell(row=r, column=c)
                val = cell.value
                cell_color = hex_from_cell(cell)
                
                if val is None or str(val).strip() == "":
                    continue

                rx = bbox_x0 + (c - 1) * col_width
                ry = bbox_y0 + (max_row - r) * row_height
                rw = col_width
                rh = row_height

                draw_styled_cell(rx, ry, rw, rh, val, cell_color)

        if colori_legenda:
            legend_elements = [Patch(facecolor=color, edgecolor='#64748b', label=label) 
                               for color, label in colori_legenda.items()]

            fig.legend(handles=legend_elements, loc='lower center', ncol=min(len(legend_elements), 4),
                       frameon=True, facecolor='#1e293b', edgecolor='#0284c7', 
                       fontsize=10, labelcolor='#f8fafc', bbox_to_anchor=(0.5, 0.02))

        nome_file_sanificato = "".join([c for c in sheet_name if c.isalnum() or c in (' ', '_')]).strip().replace(' ', '_')
        local_img = f"display_{nome_file_sanificato}.jpg"
        
        plt.savefig(local_img, facecolor=fig.get_facecolor(), edgecolor='none', bbox_inches='tight')
        print(f"✓ Immagine stanza salvata: {os.path.abspath(local_img)}")

        frame_img = os.path.join(FRAME_DRIVE, f"ditali_{nome_file_sanificato}.jpg")
        try:
            plt.savefig(frame_img, facecolor=fig.get_facecolor(), edgecolor='none', bbox_inches='tight')
            print(f"✓ Copia cornice aggiornata: {frame_img}")
        except Exception:
            pass

        plt.close()

if __name__ == "__main__":
    genera_mappe_stanze_singole()